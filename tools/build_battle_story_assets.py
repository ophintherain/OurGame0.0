"""Import boss dialogue from the second sheet without changing exploration story code."""
from pathlib import Path
import json
import re

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "聚光灯文本剧情-带剧情ID.xlsx"
OUT = ROOT / "Assets/Modules/Story/Definitions"
LIBRARY = ROOT / "Assets/Modules/Story/StoryLibrary.asset"
SCRIPT_GUID = "2914f0dd0f88b4d4c9bbd7d5ce8e8263"
GUIDS = {
    "boss1_post": "9030dc04be8445e79fc0ed741c8e72ed",
    "boss2_pre": "d742ba24b8154368a2e9969888894ebd",
    "boss2_post": "8c941646870b48dbb9ed75bb61b4f56a",
    "boss3_post": "87a9144b29264193b8541c3e2ef54a3c",
}


def quoted(value):
    return json.dumps(str(value or ""), ensure_ascii=True)


def line(kind, speaker, body, note="", scene="", portrait=""):
    return [
        f"  - kind: {kind}",
        f"    speaker: {quoted(speaker)}",
        f"    text: {quoted(body)}",
        "    choices: []",
        f"    portraitKey: {quoted(portrait)}",
        "    audioCue: \"\"",
        f"    stageNote: {quoted(note)}",
        f"    sceneTag: {quoted(scene)}",
        "    note: \"\"",
    ]


sheet = openpyxl.load_workbook(BOOK, read_only=True, data_only=True).worksheets[1]


def scene_rows(start, end, portrait=""):
    result = []
    for number in range(start, end + 1):
        note, scene, speaker, body = (sheet.cell(number, column).value for column in range(1, 5))
        if not body:
            continue
        result.append(line(1 if speaker == "旁白" else 0, speaker, body, note, scene,
                           portrait if speaker == "？？" else ""))
    return result


stories = {
    "boss1_post": scene_rows(7, 8, "黑鸟"),
    "boss2_pre": scene_rows(13, 18, "黑女孩"),
    "boss2_post": scene_rows(19, 25, "黑女孩"),
    "boss3_post": scene_rows(33, 40, "黑研究员"),
}
stories["boss2_post"].append([
    "  - kind: 5",
    "    speaker: \"\"",
    "    text: \"\"",
    "    choices:",
    f"    - {quoted(sheet.cell(26, 4).value)}",
    f"    - {quoted(sheet.cell(27, 4).value)}",
    "    portraitKey: \"\"",
    "    audioCue: \"\"",
    "    stageNote: \"\"",
    f"    sceneTag: {quoted(sheet.cell(26, 2).value)}",
    "    note: \"\"",
])
stories["boss2_post"].extend(scene_rows(28, 28))

for story_id, rows in stories.items():
    path = OUT / f"{story_id}.asset"
    content = [
        "%YAML 1.1", "%TAG !u! tag:unity3d.com,2011:", "--- !u!114 &11400000",
        "MonoBehaviour:", "  m_ObjectHideFlags: 0", "  m_CorrespondingSourceObject: {fileID: 0}",
        "  m_PrefabInstance: {fileID: 0}", "  m_PrefabAsset: {fileID: 0}",
        "  m_GameObject: {fileID: 0}", "  m_Enabled: 1", "  m_EditorHideFlags: 0",
        f"  m_Script: {{fileID: 11500000, guid: {SCRIPT_GUID}, type: 3}}",
        f"  m_Name: {story_id}", "  m_EditorClassIdentifier: ", f"  storyId: {story_id}",
        f"  chapterTitle: {story_id}", "  playOnce: 1", "  lockPlayer: 1",
        "  completeObserveTaskOnFinish: 0", "  lines:",
    ]
    for row in rows:
        content.extend(row)
    path.write_text("\n".join(content) + "\n", encoding="utf-8")
    (OUT / f"{story_id}.asset.meta").write_text(
        f"fileFormatVersion: 2\nguid: {GUIDS[story_id]}\nNativeFormatImporter:\n"
        "  externalObjects: {}\n  mainObjectFileID: 11400000\n  userData: \n"
        "  assetBundleName: \n  assetBundleVariant: \n", encoding="utf-8")

library = LIBRARY.read_text(encoding="utf-8")
for guid in GUIDS.values():
    if guid not in library:
        library += f"  - {{fileID: 11400000, guid: {guid}, type: 2}}\n"
LIBRARY.write_text(library, encoding="utf-8")
print("Imported", ", ".join(stories))
